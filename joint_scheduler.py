"""Q3: task-derived relay selection, continuous coverage certificates, joint timing.

Transport box groups/routes are a heuristic warm start from Q2. Relay locations
and all start times/resource ordering are optimized in a restricted candidate
model. No claim of unrestricted global optimality is made.
"""
import argparse
import itertools
import json
import math
from pathlib import Path
from collections import defaultdict
import numpy as np
from openpyxl import Workbook,load_workbook
from scipy.optimize import milp,Bounds,LinearConstraint
from scipy.sparse import coo_matrix
import q1_solution as q1
import q2_solution as q2
import common_physics as physics
import radio_cert as radio

terrain_link=physics.terrain_link


def xyz(node,alt=None):return node['lon'],node['lat'],node['elev'] if alt is None else alt


def communication_params(root):
    s=load_workbook(q1.find(root,'通信链路参数.xlsx'),data_only=True).active
    p={(r[0],r[3]):r[4] for r in list(s.values)[2:] if r[0] and r[3]}
    rx=p['接收参数','Psens']+p['接收参数','M'];sys=p['传播参数','Lsys']
    def budget(a,b):
        return min(p[a,'Pt'],p[b,'Pt'])+p[a,'G']+p[b,'G']-sys-rx
    return dict(direct=budget('运输无人机','固定网关 G01'),
        access=budget('运输无人机','中继接入端'),backhaul=budget('中继回传端','固定网关 G01'),
        frequency=p['传播参数','f'],obstruction=p['传播参数','Lobs'],gateway_alt=p['固定网关 G01','hG'])


def relay_data(root):
    s=load_workbook(q1.find(root,'中继无人机数据.xlsx'),data_only=True).active
    row=list(s.values)[2]
    out=dict(mass=row[4],speed=row[5],power=row[6],battery=row[7],reserve=row[8]/100,
        prep=row[9],link_time=row[10],turn=row[11],climb=row[12],descent=row[13],
        eta=row[14],hover=row[16],communication=row[17],max_height=row[18])
    inventory=list(s.values)[11]
    out.update(count=int(inventory[1]),charge_full=inventory[2],
        units=[r[0] for r in list(s.values)[6:8] if r[0]])
    return out


def relay_flight(spot,relay,nodes,dem,cache=None):
    base=xyz(nodes['O01']);results=[]
    for a,b in [(base,spot),(spot,base)]:
        d,up,down,alt=physics.leg_geometry(a,b,a[2],b[2],dem,relay=True)
        seconds=d/relay['speed']+up/relay['climb']+down/relay['descent']
        energy=relay['power']*d/relay['speed']/3600+relay['mass']*physics.GRAVITY*up/(3.6e6*relay['eta'])
        results.append((seconds,energy,alt,d,up,down))
    return results


def read_transport(path,boxes,models,nodes,dem):
    wb=load_workbook(path,data_only=True);assigned=defaultdict(list);trips=[];cache={}
    for row in list(wb['Q2_逐箱交付'].values)[1:]:
        if row[0]:assigned[row[1]].append(row[0])
    for row in list(wb['Q2_运输架次'].values)[1:]:
        if not row[0]:continue
        code,u,g,b,start,visits,end,energy=row[:8];route=visits.split('→');ids=assigned[code]
        p=q2.profile(models[g],route,ids,boxes,nodes,dem,cache)
        if p is None:raise ValueError(f'{code}: Q2路径在统一物理模型下不可行')
        if abs(p['energy']-energy)>1e-5 or abs(start+p['duration']-end)>1e-4:
            raise ValueError(f'{code}: Q2结果过期，请先重跑修正后的第二问')
        t=dict(code=code,model=g,drone=u,battery=b,ids=ids,route=route,profile=p)
        t['stages']=radio.trajectory(t,models,boxes,nodes,dem,q2.leg,cache)
        trips.append(t)
    return trips


def candidate_points(nodes,dem,relay):
    # Geographic warm starts are optional positions, never task-ID associations.
    anchors=[(109.28305555555556,23.019444444444446),
             (109.19972222222222,23.014722222222222),
             (109.23805555555556,23.070833333333333)]
    base=nodes['O01'];points=list(anchors)
    for n in nodes.values():
        points.append((n['lon'],n['lat']))
        points.append(((n['lon']+base['lon'])/2,(n['lat']+base['lat'])/2))
    lons=dem['longitude'].ravel();lats=dem['latitude'].ravel();out={}
    for x,y in points:
        col,row=physics.grid_xy((x,y),dem);col=int(round(col));row=int(round(row))
        if not (0<=row<len(lats) and 0<=col<len(lons)):continue
        z=float(dem['dem'][row,col])
        if not math.isfinite(z) or z==float(dem['nodata'].ravel()[0]):continue
        for h in [relay['max_height'],relay['max_height']/2]:
            spot=(float(lons[col]),float(lats[row]),z+h)
            if spot not in out.values():out[f'C{len(out)+1:03d}']=spot
    return out


def build_profiles(trips,nodes,dem,relay,comms):
    gateway=xyz(nodes['O01'],nodes['O01']['elev']+comms['gateway_alt'])
    candidates=candidate_points(nodes,dem,relay);valid={};flights={}
    for c,spot in candidates.items():
        cert=radio.certify_link(gateway,spot,spot,dem,comms['backhaul'],comms['frequency'],comms['obstruction'])
        if not cert:continue
        out,back=relay_flight(spot,relay,nodes,dem)
        if out[1]+back[1]>=relay['battery']*(1-relay['reserve']):continue
        valid[c]=spot;flights[c]=(out,back,cert)
    options={};direct={};caches={}
    for idx,t in enumerate(trips):
        cache={};caches[idx]=cache
        cert=radio.coverage(t['stages'],gateway,None,dem,comms,cache)
        if cert is not None:
            direct[idx]=cert;continue
        target=nodes[t['route'][-1]]
        nearest=sorted(valid,key=lambda c:physics.horizontal_distance(valid[c],xyz(target)))[:4]
        seed=[c for c in list(candidates)[:6] if c in valid]
        preferred=list(dict.fromkeys(seed+nearest));options[idx]={}
        for c in preferred:
            profile=radio.coverage(t['stages'],gateway,valid[c],dem,comms,cache)
            if profile is not None:options[idx][c]=profile
        if not options[idx]:
            for c in valid:
                if c in preferred:continue
                profile=radio.coverage(t['stages'],gateway,valid[c],dem,comms,cache)
                if profile is not None:options[idx][c]=profile
        if not options[idx]:raise RuntimeError(f"{t['code']}: 候选集中无可认证中继，需扩展位置或改路线；不代表原问题不可行")
        print('通信候选',t['code'],len(options[idx]),flush=True)
    if not options:return {},flights,options,direct
    ids=sorted({c for opts in options.values() for c in opts})
    matrix=np.array([[int(c in options[i]) for c in ids] for i in options],dtype=float)
    costs=np.array([1+1e-4*(flights[c][0][1]+flights[c][1][1]) for c in ids])
    result=milp(costs,integrality=np.ones(len(ids)),bounds=Bounds(0,1),
        constraints=LinearConstraint(matrix,1,np.inf),options={'time_limit':30})
    if result.x is None:raise RuntimeError('候选中继覆盖问题未找到可行解')
    chosen={c:valid[c] for c,v in zip(ids,result.x) if v>.5}
    if len(chosen)>relay['count']:raise RuntimeError('单次任务候选需要过多中继架次，请扩展架次复用模型')
    print('选中中继位置',list(chosen),flush=True)
    return chosen,flights,options,direct


class LinearModel:
    def __init__(self):self.lb=[];self.ub=[];self.integer=[];self.names=[];self.rows=[];self.lo=[];self.hi=[]
    def var(self,name,lb=0,ub=np.inf,binary=False):
        i=len(self.lb);self.names.append(name);self.lb.append(lb);self.ub.append(1 if binary else ub);self.integer.append(int(binary));return i
    def add(self,row,lo=-np.inf,hi=np.inf):self.rows.append(row);self.lo.append(lo);self.hi.append(hi)
    def solve(self,cost,seconds=60):
        rr=[];cc=[];vv=[]
        for i,row in enumerate(self.rows):
            for j,v in row.items():rr.append(i);cc.append(j);vv.append(v)
        A=coo_matrix((vv,(rr,cc)),shape=(len(self.rows),len(self.lb))).tocsc()
        c=np.zeros(len(self.lb))
        for i,v in cost.items():c[i]=v
        return milp(c,integrality=self.integer,bounds=Bounds(self.lb,self.ub),
            constraints=LinearConstraint(A,self.lo,self.hi),options={'time_limit':seconds,'mip_rel_gap':1e-7})


def optimize_timing(trips,spots,flights,options,direct,models,boxes,relay,batteries,horizon=40000):
    M=3*horizon;guard=.01;lp=LinearModel();starts={};late={};assign={};rs={};re={};units={}
    for i,t in enumerate(trips):
        hard=min((q2.hard_deadline(boxes[b])-d for b,d in t['profile']['deliveries'].items()),default=np.inf)
        if hard<0:raise RuntimeError('某货箱在零起始时刻也无法按时完成')
        starts[i]=lp.var(f'transport_start_{i}',ub=min(horizon,hard))
        for b,d in t['profile']['deliveries'].items():
            late[b]=lp.var('late_'+b,ub=horizon)
            lp.add({late[b]:1,starts[i]:-1},lo=d-boxes[b]['desired'])
    endall=lp.var('makespan',ub=horizon)
    for i,t in enumerate(trips):lp.add({endall:1,starts[i]:-1},lo=t['profile']['duration'])
    for c in spots:
        out,back,_=flights[c];conn=relay['prep']+out[0]+relay['link_time']
        base_energy=out[1]+back[1]+relay['hover']*relay['link_time']/3600
        span=(relay['battery']*(1-relay['reserve'])-base_energy)*3600/(relay['hover']+relay['communication'])
        rs[c]=lp.var('relay_start_'+c,ub=horizon);re[c]=lp.var('relay_service_end_'+c,ub=horizon)
        lp.add({re[c]:1,rs[c]:-1},lo=conn,hi=conn+span)
        lp.add({endall:1,re[c]:-1},lo=back[0])
        for u in relay['units']:units[c,u]=lp.var(f'relay_unit_{c}_{u}',binary=True)
        lp.add({units[c,u]:1 for u in relay['units']},lo=1,hi=1)
    for i in options:
        opts=[c for c in spots if c in options[i]]
        for c in opts:
            assign[i,c]=lp.var(f'assign_{i}_{c}',binary=True)
            needed=[p for p in options[i][c] if p['mode']=='中继']
            first=min(p['start'] for p in needed);last=max(p['end'] for p in needed)
            conn=relay['prep']+flights[c][0][0]+relay['link_time']
            lp.add({starts[i]:1,rs[c]:-1,assign[i,c]:-M},lo=conn+guard-first-M)
            lp.add({re[c]:1,starts[i]:-1,assign[i,c]:-M},lo=last+guard-M)
        lp.add({assign[i,c]:1 for c in opts},lo=1,hi=1)
    for c in spots:lp.add({v:1 for (i,site),v in assign.items() if site==c},lo=1)
    for i,j in itertools.combinations(range(len(trips)),2):
        a,b=trips[i],trips[j];same_unit=a['drone']==b['drone'];same_bat=a['battery']==b['battery']
        if not same_unit and not same_bat:continue
        order=lp.var(f'transport_order_{i}_{j}',binary=True)
        lag=[]
        for t in (a,b):
            extra=q2.charge_time(t['profile']['energy'],models[t['model']]['battery'],batteries[t['model']][0][1]) if same_bat else 0
            lag.append(t['profile']['duration']+extra+guard)
        lp.add({starts[j]:1,starts[i]:-1,order:-M},lo=lag[0]-M)
        lp.add({starts[i]:1,starts[j]:-1,order:M},lo=lag[1])
    for c,d in itertools.combinations(spots,2):
        order=lp.var(f'relay_order_{c}_{d}',binary=True)
        for u in relay['units']:
            lp.add({rs[c]:1,re[d]:-1,units[c,u]:-M,units[d,u]:-M,order:-M},lo=flights[d][1][0]+relay['turn']+guard-3*M)
            lp.add({rs[d]:1,re[c]:-1,units[c,u]:-M,units[d,u]:-M,order:M},lo=flights[c][1][0]+relay['turn']+guard-2*M)
    objective={late[b]:boxes[b]['weight'] for b in late}
    first=lp.solve(objective)
    if first.x is None:return None
    lp.add(objective,hi=max(0,float(first.fun))+1e-6)
    second=lp.solve({endall:1})
    result=second if second.x is not None else first
    timings=result.x
    if second.x is not None:
        lp.add({endall:1},hi=second.x[endall]+.01)
        energy={}
        for c in spots:energy[re[c]]=1;energy[rs[c]]=-1
        if energy:
            third=lp.solve(energy)
            if third.x is not None:timings=third.x;result=third
    missions=[];planned=[];coverage={}
    for n,c in enumerate(spots,1):
        out,back,certificate=flights[c];s=float(timings[rs[c]]);end=float(timings[re[c]])
        connected=s+relay['prep']+out[0]+relay['link_time'];finish=end+back[0]
        energy=out[1]+back[1]+relay['hover']*relay['link_time']/3600+(relay['hover']+relay['communication'])*(end-connected)/3600
        unit=next(u for u in relay['units'] if timings[units[c,u]]>.5)
        missions.append(dict(code=f'Q3-R-{n:03d}',spot=c,position=spots[c],unit=unit,component=f'R-BAT-{n:02d}',
            start=s,connected=connected,service_end=end,finish=finish,energy=energy,
            soc=100*(1-energy/relay['battery']),battery_ready=finish+q2.charge_time(energy,relay['battery'],relay['charge_full']),backhaul_certificate=certificate))
    byspot={m['spot']:m for m in missions}
    for i,t in enumerate(trips):
        start=float(timings[starts[i]]);p=t['profile']
        item=dict(t,start=start,end=start+p['duration'],energy=p['energy'],
            deliveries={b:start+d for b,d in p['deliveries'].items()},
            battery_ready=start+p['duration']+q2.charge_time(p['energy'],models[t['model']]['battery'],batteries[t['model']][0][1]))
        planned.append(item)
        if i in direct:parts=direct[i];mission=None
        else:
            c=next(c for c in spots if (i,c) in assign and timings[assign[i,c]]>.5)
            parts=options[i][c];mission=byspot[c]
        coverage[t['code']]=[dict(p,start=start+p['start'],end=start+p['end'],
            relay=mission['code'] if p['mode']=='中继' else None) for p in parts]
    status={'restricted_solver_status':int(result.status),'restricted_solver_message':result.message,
            'weighted_tardiness':sum(boxes[b]['weight']*max(0,t['deliveries'][b]-boxes[b]['desired']) for t in planned for b in t['ids']),
            'horizon_s':horizon,'model_scope':'Q2路线/装箱/运输资源ID固定；中继候选位置、任务开始及资源使用顺序联合优化'}
    return planned,missions,coverage,status


def validate_joint(trips,missions,coverage,boxes,models,fleet,batteries,relay):
    q2.validate(trips,boxes,models,fleet,batteries)
    for m in missions:
        assert m['soc']>=100*relay['reserve']-1e-6
        assert m['start']>=-1e-6 and m['connected']<=m['service_end']+1e-6
    for u in relay['units']:
        seq=sorted([m for m in missions if m['unit']==u],key=lambda m:m['start'])
        assert all(a['finish']+relay['turn']<=b['start']+1e-6 for a,b in zip(seq,seq[1:]))
    assert len({m['component'] for m in missions})==len(missions)<=relay['count']
    lookup={m['code']:m for m in missions}
    for t in trips:
        parts=coverage[t['code']]
        takeoff=t['start']+models[t['model']]['prep']+models[t['model']]['loading']*len(t['ids'])
        assert abs(parts[0]['start']-takeoff)<1e-6 and abs(parts[-1]['end']-t['end'])<1e-6
        assert all(abs(a['end']-b['start'])<1e-6 for a,b in zip(parts,parts[1:]))
        for p in parts:
            assert p['margin_db']>0
            if p['mode']=='中继':
                m=lookup[p['relay']]
                assert m['connected']<=p['start']+1e-6 and m['service_end']>=p['end']-1e-6


def output(path,trips,missions,coverage,status,boxes,models,relay,baseline):
    wb=Workbook();transport=wb.active;transport.title='Q2_运输架次'
    transport.append(['架次编号','无人机编号','机型编号','电池编号','开始时刻（s）','访问服务区顺序','返回O01时刻（s）','架次能耗（kWh）'])
    delivery=wb.create_sheet('Q2_逐箱交付');delivery.append(['货箱编号','架次编号','服务区编号','交付完成时刻（s）'])
    r=wb.create_sheet('Q3_中继架次');r.append(['中继架次编号','中继无人机编号','能源组件编号','开始时刻（s）','悬停经度（°）','悬停纬度（°）','悬停海拔（m）','建链完成时刻（s）','服务结束时刻（s）','返回O01时刻（s）','架次能耗（kWh）'])
    c=wb.create_sheet('Q3_通信保障');c.append(['运输架次编号','通信阶段','开始时刻（s）','结束时刻（s）','保障方式','中继架次编号','连续区间认证','最小保守裕量dB'])
    resources=wb.create_sheet('资源占用');resources.append(['资源类型','资源编号','架次编号','占用开始s','返回s','充电或周转后可用s'])
    due=wb.create_sheet('逐箱时限校验');due.append(['货箱编号','完成s','期望s','硬截止s','硬时限通过'])
    for t in sorted(trips,key=lambda t:(t['start'],t['code'])):
        transport.append([t['code'],t['drone'],t['model'],t['battery'],t['start'],'→'.join(t['route']),t['end'],t['energy']])
        resources.append(['运输机',t['drone'],t['code'],t['start'],t['end'],t['end']])
        resources.append(['运输电池',t['battery'],t['code'],t['start'],t['end'],t['battery_ready']])
        for b,at in sorted(t['deliveries'].items()):
            delivery.append([b,t['code'],boxes[b]['site'],at]);hd=q2.hard_deadline(boxes[b])
            due.append([b,at,boxes[b]['desired'],hd if math.isfinite(hd) else None,at<=hd+1e-6])
        merged=[]
        for p in coverage[t['code']]:
            key=(p['phase'],p['mode'],p['relay'])
            if merged and merged[-1]['key']==key:
                merged[-1]['end']=p['end'];merged[-1]['margin_db']=min(merged[-1]['margin_db'],p['margin_db'])
            else:merged.append(dict(p,key=key))
        for p in merged:c.append([t['code'],p['phase'],p['start'],p['end'],p['mode'],p['relay'],'CERTIFIED',p['margin_db']])
    for m in missions:
        x,y,z=m['position'];r.append([m['code'],m['unit'],m['component'],m['start'],x,y,z,m['connected'],m['service_end'],m['finish'],m['energy']])
        resources.append(['中继机',m['unit'],m['code'],m['start'],m['finish'],m['finish']+relay['turn']])
        resources.append(['中继组件',m['component'],m['code'],m['start'],m['finish'],m['battery_ready']])
    summary=wb.create_sheet('综合指标')
    info={'交付箱数':len(boxes),'运输架次数':len(trips),'中继架次数':len(missions),
        '联合任务完成时刻s':max([t['end'] for t in trips]+[m['finish'] for m in missions]),
        '运输能耗kWh':sum(t['energy'] for t in trips),'中继能耗kWh':sum(m['energy'] for m in missions),
        '加权迟到秒':status['weighted_tardiness'],'硬时限违规数':0,
        '通信验证':'逐线性轨迹区间的扫掠三角形/距离上界证书；不是有限采样冒充连续验证',
        '通信表语义':'列出的保障链路在整个区间可用；直连可认证时优先；中继行可保守覆盖直连与中继均可用的交叠区间',
        '求解范围':status['model_scope'],'最优性':'受限候选模型的可行优化解，不声称全问题全局最优',
        'Q2输入':str(baseline.resolve())}
    for k,v in info.items():summary.append([k,v])
    for sheet in wb:sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    path.parent.mkdir(parents=True,exist_ok=True);wb.save(path)
    proof=dict(status=status,transport=trips,relay=missions,coverage=coverage,summary=info)
    path.with_suffix('.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2,default=lambda x:float(x)),encoding='utf-8')
    print('Q3完成',info,flush=True)


def run(root,baseline,result):
    models,nodes,dem,fleet,batteries,boxes=q2.get_inputs(root)
    relay=relay_data(root);comms=communication_params(root)
    trips=read_transport(baseline,boxes,models,nodes,dem)
    spots,flights,options,direct=build_profiles(trips,nodes,dem,relay,comms)
    solution=optimize_timing(trips,spots,flights,options,direct,models,boxes,relay,batteries)
    if solution is None:
        raise RuntimeError('当前选点和固定运输候选下未找到可行时序；请扩大候选，不能解释为原题无解')
    planned,missions,coverage,status=solution
    validate_joint(planned,missions,coverage,boxes,models,fleet,batteries,relay)
    output(result,planned,missions,coverage,status,boxes,models,relay,baseline)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--data',type=Path,default=physics.DEFAULT_DATA)
    p.add_argument('--q2',type=Path,default=physics.OUTPUT/'Q2_结果.xlsx')
    p.add_argument('--out',type=Path,default=physics.OUTPUT/'Q3_结果.xlsx')
    a=p.parse_args()
    if a.data.suffix.lower()=='.zip':
        import tempfile,zipfile
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with zipfile.ZipFile(a.data) as z:
                for item in z.infolist():
                    if item.is_dir() or item.filename.endswith('\\'):continue
                    parts=item.filename.replace('\\','/').split('/')
                    if any(x in ('','..') or ':' in x for x in parts):raise ValueError('非法压缩包路径')
                    dest=root.joinpath(*parts);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(z.read(item))
            run(root,a.q2,a.out)
    else:run(a.data,a.q2,a.out)


if __name__=='__main__':main()
