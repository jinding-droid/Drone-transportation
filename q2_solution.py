#!/usr/bin/env python3
"""问题二：多点多架次可行调度；逐箱时限、无人机与电池充电校验。"""
import argparse
import itertools
import math
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

import q1_solution as q1

TOL=1e-6


def get_inputs(root):
    models,nodes,_,mat=q1.load_data(root)
    drone=load_workbook(q1.find(root,'运输无人机数据.xlsx'),data_only=True).active
    fleet=defaultdict(list)
    batteries={}
    for row in list(drone.values)[7:16]:
        if row[0] and str(row[0]).startswith('U'):
            fleet[row[1]].append(row[0])
    for row in list(drone.values)[19:]:
        if row[0] in models:
            batteries[row[0]]=[(f'{row[0]}-BAT-{i:02d}',float(row[2])) for i in range(1,int(row[1])+1)]
    sheet=load_workbook(q1.find(root,'物资需求与配送时限.xlsx'),data_only=True)['逐箱货箱清单']
    boxes={r[0]:dict(site=r[1],kind=r[2],mass=float(r[3]),volume=float(r[4]),first=r[5]=='是',
                     first_deadline=None if r[6] is None else float(r[6]),desired=float(r[7]),weight=float(r[8]))
           for r in list(sheet.values)[1:] if r[0]}
    return models,nodes,mat,fleet,batteries,boxes


def leg(nodes,mat,start,end,cache):
    if (start,end) in cache:return cache[(start,end)]
    a,b=nodes[start],nodes[end]
    data=mat['dem']; lons=mat['longitude'].ravel();lats=mat['latitude'].ravel()
    count=int(max(abs(b['lon']-a['lon'])/abs(lons[1]-lons[0]),
                  abs(b['lat']-a['lat'])/abs(lats[1]-lats[0]))*8)+2
    xx=np.linspace(a['lon'],b['lon'],count);yy=np.linspace(a['lat'],b['lat'],count)
    ix=np.clip(np.rint((xx-lons[0])/(lons[1]-lons[0])).astype(int),0,len(lons)-1)
    iy=np.clip(np.rint((lats[0]-yy)/(lats[0]-lats[1])).astype(int),0,len(lats)-1)
    vals=data[iy,ix]
    if np.any(~np.isfinite(vals)) or np.any(vals==float(mat['nodata'].ravel()[0])):
        raise ValueError(f'{start}->{end} 航段 DEM 无效')
    aalt=a['elev']+(0 if start=='O01' else 30)
    balt=b['elev']+(0 if end=='O01' else 30)
    altitude=max(float(vals.max())+50,aalt,balt)
    phi1,phi2=math.radians(a['lat']),math.radians(b['lat'])
    x=math.sin((phi2-phi1)/2)**2+math.cos(phi1)*math.cos(phi2)*math.sin(math.radians(b['lon']-a['lon'])/2)**2
    distance=2*6371008.8*math.asin(math.sqrt(x))
    cache[(start,end)]=(distance,altitude-aalt,altitude-balt,altitude)
    return cache[(start,end)]


def profile(model,route,ids,boxes,nodes,mat,cache):
    """总能耗/占用时长/每一站交付完成的相对时间。"""
    if not ids or not route or len(set(route))!=len(route):return None
    total_mass=sum(boxes[i]['mass'] for i in ids)
    volume=sum(boxes[i]['volume'] for i in ids)
    if total_mass>model['cap']+TOL or volume>model['volume']+TOL:return None
    if set(route)!={boxes[i]['site'] for i in ids}:return None
    now=model['prep']+model['loading']*len(ids)
    energy=0.;load=total_mass;last='O01';deliveries={};legs=[]
    for nxt in list(route)+['O01']:
        distance,up,down,alt=leg(nodes,mat,last,nxt,cache)
        effective=model['r0']-(model['r0']-model['rfull'])*load/model['cap']
        e=model['battery']*distance/effective+(model['empty']+load)*q1.G*up/(3.6e6*model['eta'])
        dt=distance/model['speed']+up/model['climb']+down/model['descent']
        energy+=e;now+=dt
        legs.append((last,nxt,load,dt,e,alt))
        if nxt!='O01':
            site_ids=[i for i in ids if boxes[i]['site']==nxt]
            now+=model['handoff']+model['perbox']*len(site_ids)
            for bid in site_ids:deliveries[bid]=now
            load-=sum(boxes[i]['mass'] for i in site_ids)
        last=nxt
    if energy>model['battery']*(1-model['reserve'])+TOL:return None
    return dict(energy=energy,duration=now,deliveries=deliveries,legs=legs,
                mass=total_mass,volume=volume,soc=100*(1-energy/model['battery']))


def charge_time(used,capacity,full_time):
    ending=1-used/capacity
    if ending<.9:
        return full_time*((.9-ending)*.65/.9+.35)
    return full_time*(1-ending)*.35/.1


def hard_deadline(box):
    values=[]
    if box['kind']=='医疗物资':values.append(box['desired'])
    if box['first']:values.append(box['first_deadline'])
    return min(values) if values else math.inf


def initial_tasks(boxes,models,nodes,mat,cache,start_model=None):
    used=set();tasks=[]
    # 3,600 秒的首批任务占据初始 8 架实体无人机；重载资源优先投入密集点。
    if start_model is None:
        start_model={'S001':'C','S002':'B','S006':'C','S007':'A',
                     'S010':'A','S012':'B','S013':'A','S014':'A'}
    for site,kind in start_model.items():
        urgent=[i for i,b in boxes.items() if b['site']==site and
                (b['kind']=='医疗物资' or b['first'])]
        urgent.sort(key=lambda i:(boxes[i]['kind']!='医疗物资',i))
        if len(urgent)!=len(set(urgent)):raise AssertionError(site)
        ids=urgent[:]
        # 增装同站货箱：优先医疗和期望送达早的货箱，保持起飞可行。
        candidates=[i for i,b in boxes.items() if b['site']==site and i not in ids]
        candidates.sort(key=lambda i:(boxes[i]['desired'],-boxes[i]['weight'],i))
        for i in candidates:
            if profile(models[kind],[site],ids+[i],boxes,nodes,mat,cache):ids.append(i)
        tasks.append(dict(ids=ids,route=[site],fixed=kind,phase=0))
        used.update(ids)
    return tasks,used


def remaining_tasks(boxes,used,models,nodes,mat,cache):
    sites=defaultdict(lambda:defaultdict(list))
    for bid,b in boxes.items():
        if bid not in used:sites[b['site']][(b['kind'],b['mass'],b['volume'])].append(bid)
    geometry=q1.leg_geometry(nodes,mat)
    tasks=[]
    for site,group in sorted(sites.items()):
        for trip in q1.solve_site(site,group,models,geometry[site],.2):
            tasks.append(dict(ids=trip['ids'],route=[site],fixed=None,phase=1))
    return tasks


def schedule(tasks,boxes,models,nodes,mat,cache,fleet,batteries):
    drone_free={drone:0. for drones in fleet.values() for drone in drones}
    battery_free={bid:0. for groups in batteries.values() for bid,_ in groups}
    service=[]
    def priority(t):
        deadlines=[hard_deadline(boxes[i]) for i in t['ids']]
        desired=min(boxes[i]['desired'] for i in t['ids'])
        return (t['phase'],min(min(deadlines),desired),min(deadlines),
                -sum(boxes[i]['weight'] for i in t['ids']))
    for task in sorted(tasks,key=priority):
        choices=[]
        model_kinds=[task['fixed']] if task['fixed'] else list(models)
        route_options=[task['route']]
        if len(task['route'])==2:route_options.append(task['route'][::-1])
        for kind in model_kinds:
            m=models[kind]
            for route in route_options:
                p=profile(m,route,task['ids'],boxes,nodes,mat,cache)
                if p is None:continue
                for drone in fleet[kind]:
                    for bid,charge_full in batteries[kind]:
                        start=max(drone_free[drone],battery_free[bid])
                        ends={i:start+dt for i,dt in p['deliveries'].items()}
                        violations=sum(max(0,ends[i]-hard_deadline(boxes[i])) for i in task['ids']
                                       if math.isfinite(hard_deadline(boxes[i])))
                        tardiness=sum(boxes[i]['weight']*max(0,ends[i]-boxes[i]['desired'])
                                      for i in task['ids'])
                        # 先保证硬时限；再减少当前架次的加权迟到与返航时刻。
                        key=(violations>1e-5,violations,tardiness,start+p['duration'],p['energy'],kind)
                        choices.append((key,kind,drone,bid,charge_full,start,route,p,ends))
        if not choices:raise RuntimeError('不存在物理可行的架次: '+','.join(task['ids']))
        _,kind,drone,bid,charge_full,start,route,p,ends=min(choices,key=lambda t:t[0])
        end=start+p['duration'];bat_ready=end+charge_time(p['energy'],models[kind]['battery'],charge_full)
        drone_free[drone]=end;battery_free[bid]=bat_ready
        service.append(dict(ids=task['ids'],route=route,model=kind,drone=drone,battery=bid,
                            start=start,end=end,battery_ready=bat_ready,profile=p,deliveries=ends))
    return service


def metrics(service,boxes):
    end=max(t['end'] for t in service)
    delivered={i:at for t in service for i,at in t['deliveries'].items()}
    late=sum(boxes[i]['weight']*max(0,at-boxes[i]['desired']) for i,at in delivered.items())
    hard=[(i,at,hard_deadline(boxes[i])) for i,at in delivered.items()
          if at>hard_deadline(boxes[i])+1e-5]
    return dict(late=late,hard=hard,makespan=end,energy=sum(t['profile']['energy'] for t in service),
                flights=len(service))


def refine(tasks,boxes,models,nodes,mat,cache,fleet,batteries):
    current=schedule(tasks,boxes,models,nodes,mat,cache,fleet,batteries)
    result=metrics(current,boxes)
    # 非首波任务两两合并，允许跨服务区；逐次重排无人机及共享电池。
    for step in range(12):
        best=None
        for i,j in itertools.combinations(range(len(tasks)),2):
            a,b=tasks[i],tasks[j]
            if a['phase']==0 or b['phase']==0:continue
            route=list(dict.fromkeys(a['route']+b['route']))
            if len(route)>2:continue
            merged=dict(ids=a['ids']+b['ids'],route=route,fixed=None,phase=1)
            if not any(profile(m,route,merged['ids'],boxes,nodes,mat,cache) or
                       (len(route)==2 and profile(m,route[::-1],merged['ids'],boxes,nodes,mat,cache))
                       for m in models.values()):continue
            candidate=[t for k,t in enumerate(tasks) if k not in (i,j)]+[merged]
            try:
                scheduled=schedule(candidate,boxes,models,nodes,mat,cache,fleet,batteries)
            except RuntimeError:
                continue
            score=metrics(scheduled,boxes)
            if score['hard']:continue
            # 软时限优先；架次数与能耗折算成秒权重，避免过度集中在单个架次。
            value=(score['late']/2+score['makespan']+200*score['flights']+30*score['energy'])
            old=(result['late']/2+result['makespan']+200*result['flights']+30*result['energy'])
            if value<old-1e-5 and (best is None or value<best[0]):best=(value,candidate,scheduled,score)
        if best is None:break
        _,tasks,current,result=best
        print(f'合并第 {step+1} 次: {result["flights"]} 架次，跨站架次 {sum(len(t["route"])>1 for t in current)}')
    return current


def validate(service,boxes,models,fleet,batteries):
    allids=[i for t in service for i in t['ids']]
    assert len(allids)==len(set(allids))==len(boxes)
    for t in service:
        m=models[t['model']];p=t['profile']
        assert p['energy']<=m['battery']*(1-m['reserve'])+TOL
        assert p['mass']<=m['cap']+TOL and p['volume']<=m['volume']+TOL
        assert set(t['route'])=={boxes[i]['site'] for i in t['ids']}
        for i,at in t['deliveries'].items():
            assert at<=hard_deadline(boxes[i])+TOL, f'{i} 超硬时限：{at}, {hard_deadline(boxes[i])}'
    for kind,devices in fleet.items():
        for unit in devices:
            a=sorted((t for t in service if t['drone']==unit),key=lambda t:t['start'])
            assert all(x['end']<=y['start']+TOL for x,y in zip(a,a[1:]))
        for bid,charge_full in batteries[kind]:
            a=sorted((t for t in service if t['battery']==bid),key=lambda t:t['start'])
            assert all(x['battery_ready']<=y['start']+TOL for x,y in zip(a,a[1:]))


def output(path,service,boxes,models,fleet,batteries):
    w=Workbook();trips=w.active;trips.title='Q2_运输架次'
    trips.append(['架次编号','无人机编号','机型编号','电池编号','开始时刻（s）','访问服务区顺序','返回O01时刻（s）','架次能耗（kWh）'])
    delivery=w.create_sheet('Q2_逐箱交付');delivery.append(['货箱编号','架次编号','服务区编号','交付完成时刻（s）'])
    detail=w.create_sheet('架次物理校验');detail.append(['架次编号','货箱编号列表','起飞载重kg','装载体积m³','返航SOC%','电池充满可复用时刻s','航段与每段能耗'])
    devices=w.create_sheet('资源占用校验');devices.append(['架次编号','无人机','电池','占用开始s','飞回中心s','电池充满s'])
    deadlines=w.create_sheet('时限核验');deadlines.append(['货箱编号','首批截止s','期望送达s','实际交付s','硬时限s','硬时限达标','迟到s','优先系数'])
    numbering={id(t):f'Q2-{k:03d}' for k,t in enumerate(sorted(service,key=lambda t:(t['start'],t['drone'])),1)}
    for t in sorted(service,key=lambda t:(t['start'],t['drone'])):
        code=numbering[id(t)];p=t['profile']
        trips.append([code,t['drone'],t['model'],t['battery'],t['start'],'→'.join(t['route']),t['end'],p['energy']])
        detail.append([code,','.join(t['ids']),p['mass'],p['volume'],p['soc'],t['battery_ready'],
                       '; '.join(f'{a}→{b}({load:.1f}kg,{dt:.1f}s,{e:.4f}kWh,{alt:.1f}m)' for a,b,load,dt,e,alt in p['legs'])])
        devices.append([code,t['drone'],t['battery'],t['start'],t['end'],t['battery_ready']])
        for i,at in sorted(t['deliveries'].items()):
            b=boxes[i];delivery.append([i,code,b['site'],at])
            hd=hard_deadline(b)
            deadlines.append([i,b['first_deadline'],b['desired'],at,None if math.isinf(hd) else hd,
                              '是' if at<=hd+TOL else '否',max(0,at-b['desired']),b['weight']])
    result=metrics(service,boxes)
    info=w.create_sheet('汇总说明');
    for name,val in [('总架次',result['flights']),('跨服务区架次',sum(len(t['route'])>1 for t in service)),
                     ('80箱交付',len(boxes)),('硬时限超时箱数',len(result['hard'])),
                     ('加权迟到秒',result['late']),('任务完成时刻秒',result['makespan']),
                     ('运输总能耗kWh',result['energy']),('运输无人机数量',sum(map(len,fleet.values()))),
                     ('共享电池组总数',sum(map(len,batteries.values()))),
                     ('说明','启发式可行方案；第1架次开始时刻为0，固定准备及装载包含于架次时段')]:info.append([name,val])
    for sheet in w:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    path.parent.mkdir(parents=True,exist_ok=True);w.save(path)
    print('输出',path,result,'跨服务区架次',sum(len(t['route'])>1 for t in service))


def main():
    a=argparse.ArgumentParser(description=__doc__)
    a.add_argument('--data',type=Path,default=Path('work'))
    a.add_argument('--out',type=Path,default=Path('Q2_结果.xlsx'))
    args=a.parse_args()
    if args.data.suffix.lower()=='.zip':
        import tempfile,zipfile,shutil
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)
            with zipfile.ZipFile(args.data) as archive:
                for item in archive.infolist():
                    if item.is_dir() or item.filename.endswith('\\'):continue
                    components=item.filename.replace('\\','/').split('/')
                    if '..' in components or not all(components):raise ValueError('不安全压缩包路径')
                    path=root.joinpath(*components);path.parent.mkdir(parents=True,exist_ok=True)
                    with archive.open(item) as source,path.open('wb') as target:shutil.copyfileobj(source,target)
            run(root,args.out)
    else:run(args.data,args.out)


def run(root,out):
    models,nodes,mat,fleet,batteries,boxes=get_inputs(root)
    cache={}
    assignment={'S001':'C','S002':'B','S006':'C','S007':'A',
                'S010':'A','S012':'B','S013':'A','S014':'A'}
    def build(mapping):
        first,used=initial_tasks(boxes,models,nodes,mat,cache,mapping)
        tasks=first+remaining_tasks(boxes,used,models,nodes,mat,cache)
        schedule0=schedule(tasks,boxes,models,nodes,mat,cache,fleet,batteries)
        result=metrics(schedule0,boxes)
        return tasks,result
    def rank(result):
        return (result['late']/2+result['makespan']+200*result['flights']+30*result['energy'])
    tasks,baseline=build(assignment)
    print('初始',baseline)
    # 交换两处首波机型，保持 4 A / 2 B / 2 C 的初始资源容量。
    for left,right in itertools.combinations(assignment,2):
        if assignment[left]==assignment[right]:continue
        candidate=assignment.copy();candidate[left],candidate[right]=candidate[right],candidate[left]
        try:
            candidate_tasks,candidate_result=build(candidate)
        except RuntimeError:
            continue
        if candidate_result['hard']:continue
        if rank(candidate_result)<rank(baseline)-1e-5:
            assignment,tasks,baseline=candidate,candidate_tasks,candidate_result
            print('改进首批机型',assignment,baseline)
    service=refine(tasks,boxes,models,nodes,mat,cache,fleet,batteries)
    validate(service,boxes,models,fleet,batteries)
    output(out,service,boxes,models,fleet,batteries)


if __name__=='__main__':main()
