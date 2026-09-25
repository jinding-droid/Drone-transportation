#!/usr/bin/env python3
"""第四题：继承问题三的架次与中继关系，枚举 2/3 个独立任务组的资源配置。"""
import argparse
import itertools
import math
import re
from collections import Counter,defaultdict
from pathlib import Path

from openpyxl import Workbook,load_workbook

import q1_solution as q1
import q2_solution as q2
import q3_solution as q3
import common_physics as physics

KINDS=['A无人机','B无人机','C无人机','A电池','B电池','C电池','中继无人机','中继能源组件']


def read_q3(path):
    w=load_workbook(path,data_only=True)
    transports={r[0]:dict(code=r[0],drone=r[1],model=r[2],battery=r[3],start=float(r[4]),
                          route=r[5].split('→'),end=float(r[6]),energy=float(r[7]))
                for r in list(w['Q2_运输架次'].values)[1:]}
    deliveries=defaultdict(list)
    for r in list(w['Q2_逐箱交付'].values)[1:]:deliveries[r[1]].append(r[0])
    relays={r[0]:dict(code=r[0],drone=r[1],battery=r[2],start=float(r[3]),lon=float(r[4]),
                       lat=float(r[5]),alt=float(r[6]),connected=float(r[7]),
                       service_end=float(r[8]),end=float(r[9]),energy=float(r[10]))
            for r in list(w['Q3_中继架次'].values)[1:]}
    support=defaultdict(set)
    for r in list(w['Q3_通信保障'].values)[1:]:
        if r[4]=='中继':support[r[5]].add(r[0])
    return transports,deliveries,relays,support


def components(sites,transports,relays,support):
    parent={s:s for s in sites}
    def find(a):
        if parent[a]!=a:parent[a]=find(parent[a])
        return parent[a]
    def join(a,b):parent[find(a)]=find(b)
    # 一条运输架次不能切分；同一中继架次也不能拆分到两个独立组。
    for t in transports.values():
        for s in t['route'][1:]:join(t['route'][0],s)
    for name,tripcodes in support.items():
        involved=sorted({site for code in tripcodes for site in transports[code]['route']})
        for s in involved[1:]:join(involved[0],s)
    clusters=defaultdict(list)
    for site in sorted(sites):clusters[find(site)].append(site)
    return sorted(clusters.values(),key=lambda c:(-len(c),c))


def interval_minimum(intervals):
    # 固定任务时间下，等质资源需求精确等于区间图的最大同时占用数。
    events=[]
    for start,end in intervals:
        if end<=start:raise ValueError((start,end))
        events.extend([(physics.time_tick(start),1),(physics.time_tick(end),-1)])
    used=peak=0
    for at,change in sorted(events,key=lambda x:(x[0],x[1])):
        used+=change;peak=max(peak,used)
    assert used==0
    return peak


def allocation(group,transports,deliveries,relays,support,models,battery_charges,relay_model):
    locations=set(group)
    fleet=[t for t in transports.values() if locations.intersection(t['route'])]
    assert all(set(t['route'])<=locations for t in fleet)
    relay_codes={name for name,codes in support.items()
                 if any(set(transports[code]['route']) & locations for code in codes)}
    selected=[relays[k] for k in sorted(relay_codes)]
    min_need={k:0 for k in KINDS};fixed_need={k:0 for k in KINDS}
    for kind in models:
        a=[t for t in fleet if t['model']==kind]
        min_need[kind+'无人机']=interval_minimum([(t['start'],t['end']) for t in a])
        fixed_need[kind+'无人机']=len({t['drone'] for t in a})
        intervals=[(t['start'],t['end']+q2.charge_time(t['energy'],models[kind]['battery'],battery_charges[kind])) for t in a]
        min_need[kind+'电池']=interval_minimum(intervals)
        fixed_need[kind+'电池']=len({t['battery'] for t in a})
    min_need['中继无人机']=interval_minimum([(r['start'],r['end']+relay_model['turn']) for r in selected])
    fixed_need['中继无人机']=len({r['drone'] for r in selected})
    min_need['中继能源组件']=interval_minimum([
        (r['start'],r['end']+q2.charge_time(r['energy'],relay_model['battery'],relay_model['charge_full']))
        for r in selected])
    fixed_need['中继能源组件']=len({r['battery'] for r in selected})
    for k in KINDS:assert fixed_need[k]>=min_need[k],(k,fixed_need[k],min_need[k])
    num_boxes=sum(len(deliveries[t['code']]) for t in fleet)
    work=sum(t['end']-t['start'] for t in fleet)+sum(r['end']-r['start'] for r in selected)
    return dict(sites=sorted(group),trips=len(fleet),boxes=num_boxes,relay_trips=len(selected),
                flight=sum(t['end']-t['start'] for t in fleet),work=work,
                transport_energy=sum(t['energy'] for t in fleet),relay_energy=sum(r['energy'] for r in selected),
                min_need=min_need,fixed_need=fixed_need)


def partitions(clusters,k):
    # 限定组号首次出现按 0,1,...，消除等价的组号置换。
    labels=[0]*len(clusters)
    def rec(index,maximum):
        if index==len(clusters):
            if maximum+1==k:
                groups=[[] for _ in range(k)]
                for atom,group in zip(clusters,labels):groups[group].extend(atom)
                yield groups
            return
        for value in range(min(maximum+1,k-1)+1):
            if value>maximum+1:continue
            labels[index]=value
            yield from rec(index+1,max(maximum,value))
    yield from rec(1,0)


def compute_all(clusters,k,transports,deliveries,relays,support,models,charges,relay,stock):
    results=[]
    for groups in partitions(clusters,k):
        data=[allocation(g,transports,deliveries,relays,support,models,charges,relay) for g in groups]
        assert sum(r['boxes'] for r in data)==sum(map(len,deliveries.values()))
        total={key:sum(d['min_need'][key] for d in data) for key in KINDS}
        fixed={key:sum(d['fixed_need'][key] for d in data) for key in KINDS}
        deficit={key:max(0,total[key]-stock[key]) for key in KINDS}
        fixed_deficit={key:max(0,fixed[key]-stock[key]) for key in KINDS}
        works=[d['work'] for d in data]
        mean=sum(works)/len(works)
        cv=(sum((v-mean)**2 for v in works)/len(works))**.5/mean
        boxes=[d['boxes'] for d in data]
        bmean=sum(boxes)/len(boxes)
        bcv=(sum((v-bmean)**2 for v in boxes)/len(boxes))**.5/bmean
        score=(sum(deficit.values()),sum(fixed_deficit.values()),round(cv,8),
               round(bcv,8),sum(total.values()))
        results.append(dict(groups=data,min_total=total,fixed_total=fixed,
                            deficit=deficit,fixed_deficit=fixed_deficit,cv=cv,box_cv=bcv,score=score))
    return sorted(results,key=lambda x:x['score'])


def make_book(path,clusters,chosen,alternatives,stock):
    w=Workbook();s=w.active;s.title='Q4_分区配置'
    s.append(['K（2或3）','任务组编号','服务区列表','A型运输无人机数','B型运输无人机数','C型运输无人机数',
              'A型电池组数','B型电池组数','C型电池组数','中继无人机数','中继能源组件数'])
    metrics=w.create_sheet('分组工作量');metrics.append(['分区数','组号','服务区','货箱数','运输架次','中继架次',
                                         '运输累计占用秒','运输加中继累计占用秒','运输能耗kWh','中继能耗kWh'])
    total=w.create_sheet('库存_缺口_冗余');total.append(['分区数','资源','现有库存','最小配置合计','最小配置缺口',
                                           '最小配置库存余量','沿用Q3设备编号合计','保留原设备编号时缺口'])
    summary=w.create_sheet('备选方案比较');summary.append(['分区数','服务区划分','最小资源缺口合计',
                                           '保留设备编号缺口合计','工作量变异系数','箱数变异系数',
                                           '最小配置总资源数'])
    atoms=w.create_sheet('必须同组的服务区');atoms.append(['约束连通块','服务区列表'])
    for i,a in enumerate(clusters,1):atoms.append([i,','.join(a)])
    for k in (2,3):
        item=chosen[k]
        for i,g in enumerate(item['groups'],1):
            s.append([k,f'G{i}',','.join(g['sites']),*[g['min_need'][x] for x in KINDS]])
            metrics.append([k,f'G{i}',','.join(g['sites']),g['boxes'],g['trips'],g['relay_trips'],
                            g['flight'],g['work'],g['transport_energy'],g['relay_energy']])
        for key in KINDS:
            need=item['min_total'][key];fixed=item['fixed_total'][key]
            total.append([k,key,stock[key],need,item['deficit'][key],max(0,stock[key]-need),
                          fixed,item['fixed_deficit'][key]])
        for alt in alternatives[k]:
            summary.append([k,' | '.join(','.join(g['sites']) for g in alt['groups']),
                            sum(alt['deficit'].values()),sum(alt['fixed_deficit'].values()),
                            alt['cv'],alt['box_cv'],sum(alt['min_total'].values())])
    for sheet in w:sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    path.parent.mkdir(parents=True,exist_ok=True);w.save(path)


def run(root,q3path,out):
    models,nodes,dem,fleet,batteries,boxes=q2.get_inputs(root)
    transports,deliveries,relays,support=read_q3(q3path)
    clusters=components(sorted({b['site'] for b in boxes.values()}),transports,relays,support)
    relay=q3.relay_data(root)
    drone=load_workbook(q1.find(root,'运输无人机数据.xlsx'),data_only=True).active
    charges={row[0]:float(row[2]) for row in drone.values if row[0] in models and isinstance(row[2],(int,float)) and row[1] in (6,4)}
    stock={kind+'无人机':len(fleet[kind]) for kind in models}
    stock.update({kind+'电池':len(batteries[kind]) for kind in models})
    stock['中继无人机']=2;stock['中继能源组件']=relay['count']
    if len(clusters)<3:raise RuntimeError('严格保持中继任务对应关系时不足以划为三组')
    candidates={k:compute_all(clusters,k,transports,deliveries,relays,support,
                              models,charges,relay,stock) for k in (2,3)}
    # 主方案设立工作量不均衡上限，然后在可接受范围内减少资源缺口。
    # 全部备选（包括零阈值的资源优先方案）另列于“备选方案比较”。
    cv_limit={2:.30,3:.75}
    chosen={}
    for k,values in candidates.items():
        balanced=[v for v in values if v['cv']<=cv_limit[k]]
        chosen[k]=min(balanced or values,key=lambda v:v['score'])
    make_book(out,clusters,chosen,candidates,stock)
    print('不可拆连通块：',clusters)
    for k in (2,3):
        a=chosen[k]
        print('K=',k,'组',[(g['sites'],g['boxes'],g['trips'],g['relay_trips']) for g in a['groups']],
              '最小配置',a['min_total'],'缺口',a['deficit'],
              '固定编号配置',a['fixed_total'],'固定编号缺口',a['fixed_deficit'],
              '工作量CV',round(a['cv'],4),'备选',len(candidates[k]))


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--data',type=Path,default=physics.DEFAULT_DATA)
    p.add_argument('--q3',type=Path,default=physics.OUTPUT/'Q3_结果.xlsx')
    p.add_argument('--out',type=Path,default=physics.OUTPUT/'Q4_结果.xlsx')
    args=p.parse_args()
    if args.data.suffix.lower()=='.zip':
        import tempfile,zipfile,shutil
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with zipfile.ZipFile(args.data) as z:
                for item in z.infolist():
                    if item.is_dir() or item.filename.endswith('\\'):continue
                    fields=item.filename.replace('\\','/').split('/')
                    if '..' in fields or not all(fields):raise ValueError('不安全压缩包路径')
                    path=root.joinpath(*fields);path.parent.mkdir(parents=True,exist_ok=True)
                    with z.open(item) as source,path.open('wb') as dest:shutil.copyfileobj(source,dest)
            run(root,args.q3,args.out)
    else:run(args.data,args.q3,args.out)


if __name__=='__main__':main()
