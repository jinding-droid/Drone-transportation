#!/usr/bin/env python3
"""问题一：DEM 航段、单点载荷、逐箱整数组批和安全余量敏感性。"""
import argparse
import itertools
import math
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

import numpy as np
from scipy.io import loadmat
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csr_matrix
from openpyxl import Workbook, load_workbook

import common_physics as physics

G = physics.GRAVITY
EPS = 1e-8


def find(root, name):
    matches = list(root.rglob(name))
    if len(matches) != 1:
        raise FileNotFoundError(f"需要恰好一个 {name}，找到 {len(matches)} 个：{root}")
    return matches[0]


def load_data(root):
    drone = load_workbook(find(root, '运输无人机数据.xlsx'), data_only=True).active
    models = {}
    for row in list(drone.values)[2:5]:
        k = row[0]
        models[k] = dict(empty=row[2], cap=row[3], volume=row[4], speed=row[5],
                         r0=row[6], rfull=row[7], battery=row[8], reserve=row[9]/100,
                         prep=row[10], loading=row[11], handoff=row[12], perbox=row[13],
                         climb=row[14], descent=row[15], eta=row[16])
    nodes_sheet = load_workbook(find(root, '调度中心与服务区.xlsx'), data_only=True).active
    nodes = {}
    for row in nodes_sheet.values:
        if isinstance(row[0], str) and (row[0] == 'O01' or row[0].startswith('S0')):
            nodes[row[0]] = dict(lon=float(row[2]), lat=float(row[3]), elev=float(row[4]))
    boxes_sheet = load_workbook(find(root, '物资需求与配送时限.xlsx'), data_only=True)['逐箱货箱清单']
    boxes = defaultdict(lambda: defaultdict(list))
    for row in list(boxes_sheet.values)[1:]:
        if row[0]:
            boxes[row[1]][(row[2], float(row[3]), float(row[4]))].append(row[0])
    mat = loadmat(find(root, '*DEM.mat'))
    return models, nodes, boxes, mat


def leg_geometry(nodes, dem):
    center = nodes['O01']
    geometry = {}
    for sid, dest in nodes.items():
        if sid == 'O01':
            continue
        a = (center['lon'], center['lat'])
        b = (dest['lon'], dest['lat'])
        d, up, down, altitude = physics.leg_geometry(a,b,center['elev'],dest['elev']+30,dem)
        geometry[sid] = (d,up,down,down,up,altitude)
    return geometry


def route(model, geom, payload):
    d, up1, down1, up2, down2, _ = geom
    effective = physics.effective_range(model, payload)
    energy = model['battery']*d/effective + (model['empty']+payload)*G*up1/(3.6e6*model['eta'])
    energy += model['battery']*d/model['r0'] + model['empty']*G*up2/(3.6e6*model['eta'])
    flight = 2*d/model['speed']+(up1+up2)/model['climb']+(down1+down2)/model['descent']
    return energy, flight


def max_payload(model, geom, reserve):
    available = model['battery']*(1-reserve)
    if route(model, geom, 0)[0] > available+EPS:
        return None
    lo, hi = 0., float(model['cap'])
    for _ in range(48):
        mid = (lo+hi)/2
        if route(model, geom, mid)[0] <= available+EPS:
            lo = mid
        else:
            hi = mid
    return lo


def solve_site(sid, groups, models, geom, reserve):
    types = sorted(groups)
    counts = [len(groups[t]) for t in types]
    options = []
    for kind, model in models.items():
        safe = max_payload(model, geom, reserve)
        if safe is None:
            continue
        for qty in itertools.product(*(range(c+1) for c in counts)):
            if not any(qty):
                continue
            mass = sum(q*t[1] for q,t in zip(qty,types))
            volume = sum(q*t[2] for q,t in zip(qty,types))
            if mass > min(safe,model['cap'])+EPS or volume > model['volume']+EPS:
                continue
            energy, flight = route(model, geom, mass)
            seconds = model['prep']+model['loading']*sum(qty)+flight+model['handoff']+model['perbox']*sum(qty)
            options.append((kind, qty, mass, volume, energy, seconds))
    if not options:
        raise RuntimeError(f'{sid} 无可行组批')
    matrix = csr_matrix(np.asarray([o[1] for o in options], dtype=float).T)
    constraint = LinearConstraint(matrix, counts, counts)
    bounds = Bounds(np.zeros(len(options)), np.full(len(options), np.inf))
    integrality = np.ones(len(options))
    def optimize(costs, extra=()):
        result = milp(np.asarray(costs), integrality=integrality, bounds=bounds,
                      constraints=[constraint,*extra], options={'time_limit':120,'mip_rel_gap':0.00001})
        if result.x is None or result.status != 0:
            raise RuntimeError(f'{sid} MILP 未获最优解：{result.message}')
        return result
    ones = np.ones(len(options))
    first = optimize(ones)
    n_flights = int(round(first.fun))
    flight_constraint = LinearConstraint(ones.reshape(1,-1),n_flights,n_flights)
    energies = np.array([o[4] for o in options])
    second = optimize(energies,[flight_constraint])
    energy_constraint = LinearConstraint(energies.reshape(1,-1),0,second.fun+1e-7)
    third = optimize([o[5] for o in options],[flight_constraint,energy_constraint])
    remaining = {t:list(groups[t]) for t in types}
    trips = []
    for option, repetitions in zip(options, np.rint(third.x).astype(int)):
        kind, qty, mass, volume, energy, seconds = option
        for _ in range(repetitions):
            ids=[]
            for t,q in zip(types,qty):
                ids.extend(remaining[t][:q]); del remaining[t][:q]
            trips.append(dict(site=sid,model=kind,ids=ids,mass=mass,volume=volume,
                              energy=energy,seconds=seconds,soc=100*(1-energy/models[kind]['battery'])))
    assert all(not remaining[t] for t in types)
    return trips


def compute(models,nodes,boxes,geometry,reserve):
    trips=[]
    for sid in sorted(boxes):
        trips.extend(solve_site(sid,boxes[sid],models,geometry[sid],reserve))
    ids=[box for trip in trips for box in trip['ids']]
    assert len(ids)==len(set(ids))==sum(len(g) for site in boxes.values() for g in site.values())
    return trips


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data',type=Path,default=physics.DEFAULT_DATA,help='解压后的资料根目录')
    parser.add_argument('--out',type=Path,default=physics.OUTPUT/'Q1_结果.xlsx')
    parser.add_argument('--reserves',type=float,nargs='+',default=[0.1,0.2,0.3],help='敏感性分析余量比例，例如 0.1 0.2 0.3')
    args=parser.parse_args()
    temp=None
    if args.data.is_file() and args.data.suffix.lower()=='.zip':
        temp=tempfile.TemporaryDirectory()
        root=Path(temp.name)
        with zipfile.ZipFile(args.data) as archive:
            for item in archive.infolist():
                if item.is_dir() or item.filename.endswith('\\'):
                    continue
                components=item.filename.replace('\\','/').split('/')
                if '..' in components or any(not x for x in components):
                    raise ValueError(f'不安全的压缩包路径: {item.filename}')
                target=root.joinpath(*components)
                target.parent.mkdir(parents=True,exist_ok=True)
                with archive.open(item) as source, target.open('wb') as output:
                    import shutil
                    shutil.copyfileobj(source,output)
    else:
        root=args.data
    models,nodes,boxes,dem=load_data(root)
    geometry=leg_geometry(nodes,dem)
    wb=Workbook(); sheet=wb.active;sheet.title='Q1_单点组批'
    sheet.append(['架次编号','服务区编号','机型编号','货箱编号列表','总质量（kg）','总体积（m³）','往返时间（s）','架次能耗（kWh）','返航SOC（%）'])
    loads=wb.create_sheet('机型最大安全载荷');loads.append(['返航安全余量','服务区编号','机型','最大安全载荷（kg）','DEM最高高程上方巡航海拔（m）'])
    summary=wb.create_sheet('余量敏感性');summary.append(['返航安全余量','可行性','架次数','运输能耗（kWh）','累计作业时间（s）','A架次','B架次','C架次'])
    chosen=None
    for reserve in sorted(set(args.reserves+[0.2])):
        for sid in sorted(boxes):
            for kind,model in models.items():
                loads.append([reserve,sid,kind,max_payload(model,geometry[sid],reserve),geometry[sid][-1]])
        try:
            trips=compute(models,nodes,boxes,geometry,reserve)
            summary.append([reserve,'可行',len(trips),sum(t['energy'] for t in trips),sum(t['seconds'] for t in trips),*[sum(t['model']==k for t in trips) for k in models]])
            if abs(reserve-0.2)<1e-8:chosen=trips
        except RuntimeError as exc:
            summary.append([reserve,str(exc)])
    if chosen is None:
        raise RuntimeError('题目原始 20% 安全余量下不可行，请检查数据和模型')
    for i,t in enumerate(chosen,1):
        sheet.append([f'Q1-{i:03d}',t['site'],t['model'],','.join(t['ids']),t['mass'],t['volume'],t['seconds'],t['energy'],t['soc']])
    for s in wb:
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    args.out.parent.mkdir(parents=True,exist_ok=True)
    wb.save(args.out)
    print(f'已生成 {args.out}；架次 {len(chosen)}；总能耗 {sum(t["energy"] for t in chosen):.4f} kWh；累计作业时间 {sum(t["seconds"] for t in chosen):.1f} s')

if __name__=='__main__':
    main()
